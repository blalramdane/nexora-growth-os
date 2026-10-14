"""Safe, dry-run-first importer for NEXORA lead workbooks.

This module deliberately does not call a CRM API or mutate any CRM.
It reads the workbook, normalizes records, flags duplicate/invalid rows, and
writes a JSON report for review before any future integration is authorized.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

REQUIRED = ("المنطقة", "اسم النشاط")
PIPELINE = {
    "Lead جديد": "Lead جديد",
    "تم التواصل": "تم التواصل",
    "وافق على التجربة": "وافق على التجربة",
    "جاري تجهيز Demo": "جاري تجهيز Demo",
    "Demo اتبعت": "Demo اتبعت",
    "متابعة": "متابعة",
    "عميل متعاقد": "عميل متعاقد",
    "لم يكتمل": "لم يكتمل",
}
FIELDS = {
    "region": "المنطقة",
    "business_name": "اسم النشاط",
    "category": "التخصص",
    "address": "العنوان",
    "public_phone": "الهاتف",
    "published_rating": "التقييم المنشور",
    "published_reviews": "عدد المراجعات المنشور",
    "source_label": "مصدر البيانات",
    "qr_status": "حالة QR",
    "website_status": "حالة Website",
    "pos_status": "حالة POS",
    "lead_score": "Lead Score",
    "stage": "مرحلة CRM",
    "next_action": "الخطوة التالية",
    "last_verified_at": "تاريخ التحقق",
    "discovery_url": "مصدر الاكتشاف",
    "contact_channel": "قناة التواصل",
    "decision_maker": "الشخص المسؤول / Decision Maker",
    "verification_status": "حالة التحقق الميداني",
    "digital_signal": "إشارة الحضور الرقمي",
    "pain_signal": "Pain Signal (مؤكد/فرضية)",
    "proposed_offer": "العرض المقترح",
    "last_contact_at": "تاريخ آخر تواصل",
    "contact_outcome": "نتيجة التواصل",
    "follow_up_at": "موعد المتابعة",
    "verification_notes": "ملاحظات التحقق",
}

def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()

def normalize_text(value: Any) -> str:
    text = unicodedata.normalize("NFKC", clean(value)).casefold()
    text = re.sub(r"[\u064B-\u065F\u0670\u0640]", "", text)
    text = text.translate(str.maketrans({"أ":"ا","إ":"ا","آ":"ا","ى":"ي","ة":"ه"}))
    return re.sub(r"[^\w]+", "", text, flags=re.UNICODE)

def normalize_phone(value: Any) -> str:
    """Preserve digits as strings and leading zeros; do not guess country codes."""
    raw = clean(value)
    digits = "".join(ch for ch in raw if ch.isdigit())
    return digits

def identity_key(record: dict[str, Any]) -> tuple[str, str, str]:
    phone = normalize_phone(record.get("public_phone"))
    name = normalize_text(record.get("business_name"))
    region = normalize_text(record.get("region"))
    # Use both phone and business+region; blank phone must not collapse unrelated leads.
    return (phone, name, region)

def map_row(row: dict[str, Any], row_number: int) -> dict[str, Any]:
    missing = [key for key in REQUIRED if not clean(row.get(key))]
    if missing:
        raise ValueError(f"row {row_number}: missing required field(s): {', '.join(missing)}")
    out = {target: clean(row.get(source)) for target, source in FIELDS.items()}
    out["public_phone"] = clean(row.get("الهاتف"))  # strings preserve leading zero
    stage = out.get("stage", "")
    if stage and stage not in PIPELINE:
        out["stage_mapping_warning"] = f"Unknown pipeline stage: {stage}"
    out["source_row"] = row_number
    out["verification_status"] = out.get("verification_status") or "غير متحقق"
    for key in ("qr_status", "website_status", "pos_status"):
        out[key] = out.get(key) or "غير معروف"
    out["is_verified"] = False  # source listing alone never verifies a lead
    return out

def read_rows(path: Path, sheet_name: str = "Restaurant Leads") -> list[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        import csv
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("XLSX support requires openpyxl: pip install openpyxl") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}. Available: {workbook.sheetnames}")
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(iterator, ())]
        return [dict(zip(headers, values)) for values in iterator if any(v is not None for v in values)]
    raise ValueError("Input must be .xlsx, .xlsm, or .csv")

def build_report(rows: list[dict[str, Any]]) -> dict[str, Any]:
    report: dict[str, Any] = {
        "mode": "DRY_RUN_ONLY",
        "mutations_performed": 0,
        "input_rows": len(rows),
        "valid_rows": 0,
        "invalid_rows": [],
        "duplicates_within_file": [],
        "records": [],
        "warnings": [
            "No CRM API is called; no records are created or changed.",
            "Public directory data is not considered verified by default.",
            "Duplicate candidates require human review before merge.",
        ],
    }
    seen: dict[tuple[str, str, str], int] = {}
    for i, row in enumerate(rows, start=2):
        try:
            record = map_row(row, i)
        except ValueError as exc:
            report["invalid_rows"].append({"row": i, "error": str(exc)})
            continue
        key = identity_key(record)
        # Flag exact business+region duplicate, or same non-empty phone + region.
        duplicate_of = None
        for old_key, old_row in seen.items():
            same_name_region = key[1] and key[1] == old_key[1] and key[2] == old_key[2]
            same_phone_region = key[0] and key[0] == old_key[0] and key[2] == old_key[2]
            if same_name_region or same_phone_region:
                duplicate_of = old_row
                break
        if duplicate_of is not None:
            report["duplicates_within_file"].append({
                "row": i, "duplicate_of_row": duplicate_of,
                "business_name": record["business_name"],
                "reason": "same normalized business+region or same non-empty phone+region",
            })
            continue
        seen[key] = i
        report["records"].append(record)
        report["valid_rows"] += 1
    report["summary"] = {
        "create_candidates": report["valid_rows"],
        "update_candidates": 0,
        "skipped_duplicates": len(report["duplicates_within_file"]),
        "invalid": len(report["invalid_rows"]),
    }
    return report

def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="NEXORA CRM import dry-run; never writes to CRM.")
    parser.add_argument("input", type=Path, help="Path to XLSX/XLSM/CSV")
    parser.add_argument("--sheet", default="Restaurant Leads")
    parser.add_argument("--report", type=Path, default=Path("crm-import-dry-run.json"))
    args = parser.parse_args(argv)
    try:
        rows = read_rows(args.input, args.sheet)
        report = build_report(rows)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    print(json.dumps({"report": str(args.report), **report["summary"], "mutations_performed": 0}, ensure_ascii=False, indent=2))
    return 0 if not report["invalid_rows"] else 1

if __name__ == "__main__":
    raise SystemExit(main())
